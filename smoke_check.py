# setup_finish.ps1 — idempotent repo + smoke-check setup
Set-StrictMode -Version Latest
$root = "C:\Users\jtahm\FlightDeck"
Push-Location $root
try {
    $smokePath = Join-Path $root "Scripts\smoke_check.py"
    $smokeBody = @'
from openpyxl import load_workbook
wb = load_workbook(r"C:\Users\jtahm\FlightDeck\Output\FlightDeck_Signals.xlsx", data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    rows = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True) if any(_))
    print(f"{name}: {rows} data rows")
'@
    if (-not (Test-Path $smokePath) -or (Get-Content $smokePath -Raw) -ne $smokeBody) {
        Set-Content -Path $smokePath -Value $smokeBody -Force
        Write-Host "smoke_check.py created/updated."
    } else {
        Write-Host "smoke_check.py present and unchanged."
    }

    $gitignore = @'Output/
*.xlsx
*.xlsm
Passenger_List.csv
.env
secrets/
__pycache__/
*.pyc
.vscode/
.venv/
'@
    $gitignorePath = Join-Path $root ".gitignore"
    if (-not (Test-Path $gitignorePath) -or (Get-Content $gitignorePath -Raw) -ne $gitignore) {
        Set-Content -Path $gitignorePath -Value $gitignore -Force
        Write-Host ".gitignore created/updated."
    } else {
        Write-Host ".gitignore present and unchanged."
    }

    $isRepo = $false
    try { git rev-parse --is-inside-work-tree 2>$null | Out-Null; if ($LASTEXITCODE -eq 0) { $isRepo = $true } } catch { $isRepo = $false }

    if (-not $isRepo) {
        Write-Host "Initializing local git repository..."
        git init | Out-Null
        git config user.name "Your Name"
        git config user.email "you@example.com"
        git add .gitignore
        $toAdd = @("Scripts\generate_workbook.py","Scripts\fd_ensure_headers.py","Scripts\smoke_check.py")
        foreach ($f in $toAdd) { if (Test-Path (Join-Path $root $f)) { git add $f } }
        git commit -m "chore: add headers to generator output; add smoke check; add .gitignore" | Out-Null
        git branch -M main | Out-Null
        Write-Host "Local git repository initialized and commit created."
    } else {
        Write-Host "Directory already a git repository; no init performed."
        Write-Host "Git status (porcelain):"
        git status --porcelain
    }

    Write-Host "`nRunning smoke_check.py..."
    python $smokePath

    Write-Host "`nFinal git summary:"
    git log --oneline -n 5
    git status --porcelain

    Write-Host "`nDone. No remote added. Provide a remote URL when you want to push."
}
finally {
    Pop-Location
}