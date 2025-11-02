from openpyxl import load_workbook
rv = r"C:\Users\jtahm\FlightDeck\Output\FlightDeck_Signals.xlsx"
wb = load_workbook(rv, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    rows = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True) if any(_))
    print(f"{name}: {rows} data rows")
