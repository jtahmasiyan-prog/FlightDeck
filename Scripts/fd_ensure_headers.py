# fd_ensure_headers.py
from openpyxl import load_workbook
from pathlib import Path
import sys

# Script: fd_ensure_headers.py
# Purpose: Ensure header rows exist in the generated FlightDeck_Signals.xlsx workbook.
# Save to: C:\Users\jtahm\FlightDeck\Scripts\fd_ensure_headers.py

OUTPUT_PATH = Path(r"C:\Users\jtahm\FlightDeck\Output\FlightDeck_Signals.xlsx")

def ensure_header(ws, header):
    first_row = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    if first_row != header:
        ws.insert_rows(1)
        for col_index, value in enumerate(header, start=1):
            ws.cell(row=1, column=col_index, value=value)

def main(path: Path):
    if not path.exists():
        print(f"ERROR: Workbook not found at {path}")
        sys.exit(1)

    wb = load_workbook(path)
    ensure_header(wb["Runway_Clearance"], ["Ticker", "Close", "MA50", "Signal"])
    ensure_header(wb["Holding_Pattern"], ["Ticker", "Close", "MA50", "Signal"])
    ensure_header(wb["Tower_Log"], ["Ticker", "Notes"])
    wb.save(path)
    print(f"Headers ensured in {path}")

if __name__ == "__main__":
    main(OUTPUT_PATH)